from __future__ import annotations

import asyncio
import io
import re
import uuid
import zipfile

from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import TokenData, get_current_user
from app.db.session import get_db
from sqlalchemy import select

from app.models.user import User
from app.schemas.video_source import (
    TagRead,
    VideoSourceCreate,
    VideoSourceListItem,
    VideoSourceListResponse,
    VideoSourceParseRequest,
    VideoSourceParseResult,
    VideoSourceRead,
    VideoSourceStatHistoryResponse,
    VideoSourceStatsResponse,
    VideoSourceTagsUpdate,
)
from app.services.video_source_service import (
    create_video_source,
    delete_video_source,
    get_stats_history,
    get_video_source_or_404,
    get_video_source_stats,
    list_video_sources,
    parse_video_url,
    replace_video_source_tags,
    trigger_download_and_upload,
)

router = APIRouter(prefix="/video-sources", tags=["video-sources"])


def _get_owner_id(current_user: TokenData = Depends(get_current_user)) -> uuid.UUID | None:
    """Query filter: admin→None (no filter), user→user_id"""
    return None if current_user.is_admin else current_user.user_id


def _get_creator_id(current_user: TokenData = Depends(get_current_user)) -> uuid.UUID:
    """Create records: always returns actual user_id"""
    return current_user.user_id


# /stats and /parse MUST be before /{vs_id} to avoid UUID matching them
@router.get("/stats", response_model=VideoSourceStatsResponse)
async def get_stats_endpoint(
    platform: str | None = Query(None),
    blogger_name: str | None = Query(None),
    tiktok_blogger_id: uuid.UUID | None = Query(None),
    tag_ids: str | None = Query(None),
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceStatsResponse:
    parsed_tag_ids: list[uuid.UUID] = []
    if tag_ids:
        parsed_tag_ids = [uuid.UUID(part.strip()) for part in tag_ids.split(",") if part.strip()]
    stats = await get_video_source_stats(
        session,
        owner_id,
        platform=platform,
        blogger_name=blogger_name,
        tiktok_blogger_id=tiktok_blogger_id,
        tag_ids=parsed_tag_ids,
    )
    return VideoSourceStatsResponse(**stats)


@router.post("/parse", response_model=VideoSourceParseResult)
async def parse_video_endpoint(
    payload: VideoSourceParseRequest,
    creator_id: uuid.UUID = Depends(_get_creator_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceParseResult:
    """Parse a video URL via yt-dlp without saving to database.
    If source_url already exists for this user, returns existing_id in response."""
    return await parse_video_url(payload.source_url, session=session, owner_id=creator_id)


async def _vs_to_read(session: AsyncSession, vs: object) -> VideoSourceRead:
    """Build VideoSourceRead, querying tags from VideoSourceTag table."""
    from app.models.tag import Tag, VideoSourceTag
    from app.schemas.tiktok_blogger import TiktokBloggerRead
    # GCS 视频链按需续签（CDN 跳过），覆写到 vs 的 local_*_url 字段
    from app.utils.gcs_signing import ensure_video_source_signed_urls
    await ensure_video_source_signed_urls(session, vs)
    tags_stmt = (
        select(Tag)
        .join(VideoSourceTag, VideoSourceTag.tag_id == Tag.id)
        .where(VideoSourceTag.video_source_id == vs.id)  # type: ignore[attr-defined]
    )
    tag_rows = (await session.execute(tags_stmt)).scalars().all()
    tags = [TagRead.model_validate(t) for t in tag_rows]
    tiktok_blogger = None
    if vs.tiktok_blogger:  # type: ignore[attr-defined]
        tiktok_blogger = TiktokBloggerRead(
            **{k: getattr(vs.tiktok_blogger, k) for k in TiktokBloggerRead.model_fields if k != "video_count" and hasattr(vs.tiktok_blogger, k)},
            video_count=0,
        )
    data = {k: getattr(vs, k) for k in VideoSourceRead.model_fields if k not in ("tags", "tiktok_blogger") and hasattr(vs, k)}
    return VideoSourceRead(**data, tags=tags, tiktok_blogger=tiktok_blogger)


@router.post("", response_model=VideoSourceRead)
async def create_video_source_endpoint(
    payload: VideoSourceCreate,
    response: Response,
    creator_id: uuid.UUID = Depends(_get_creator_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceRead:
    is_new, vs = await create_video_source(session, payload, creator_id)
    response.status_code = 201 if is_new else 200
    return await _vs_to_read(session, vs)


@router.get("", response_model=VideoSourceListResponse)
async def list_video_sources_endpoint(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=500),
    platform: str | None = Query(None),
    blogger_name: str | None = Query(None),
    tiktok_blogger_id: uuid.UUID | None = Query(None),
    tag_ids: str | None = Query(None),
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceListResponse:
    parsed_tag_ids: list[uuid.UUID] = []
    if tag_ids:
        parsed_tag_ids = [uuid.UUID(part.strip()) for part in tag_ids.split(",") if part.strip()]
    rows, total = await list_video_sources(
        session, page=page, page_size=page_size, owner_id=owner_id,
        platform=platform, blogger_name=blogger_name, tiktok_blogger_id=tiktok_blogger_id, tag_ids=parsed_tag_ids,
    )

    # GCS 视频链按需续签（CDN 跳过）。返回前覆写 local_*_url 字段
    from app.utils.gcs_signing import ensure_video_sources_signed_urls
    await ensure_video_sources_signed_urls(session, rows)

    # 批量查询创建者用户名
    owner_ids = list({r.owner_id for r in rows if r.owner_id is not None})
    username_map: dict[uuid.UUID, str] = {}
    if owner_ids:
        users = (await session.execute(select(User).where(User.id.in_(owner_ids)))).scalars().all()
        for u in users:
            username_map[u.id] = u.display_name or u.username

    # 批量查询每个视频的标签
    from app.models.tag import Tag, VideoSourceTag
    from app.schemas.tiktok_blogger import TiktokBloggerRead
    vs_ids = [r.id for r in rows]
    tags_map: dict[uuid.UUID, list] = {r.id: [] for r in rows}
    if vs_ids:
        tags_stmt = (
            select(VideoSourceTag.video_source_id, Tag)
            .join(Tag, Tag.id == VideoSourceTag.tag_id)
            .where(VideoSourceTag.video_source_id.in_(vs_ids))
        )
        for vs_id, tag in (await session.execute(tags_stmt)).all():
            tags_map[vs_id].append(TagRead.model_validate(tag))

    items = [
        VideoSourceListItem(
            **{k: getattr(r, k) for k in VideoSourceListItem.model_fields if k not in ("owner_username", "tags", "tiktok_blogger") and hasattr(r, k)},
            owner_username=username_map.get(r.owner_id) if r.owner_id else None,
            tags=tags_map.get(r.id, []),
            tiktok_blogger=TiktokBloggerRead(
                **{k: getattr(r.tiktok_blogger, k) for k in TiktokBloggerRead.model_fields if k != "video_count" and hasattr(r.tiktok_blogger, k)},
                video_count=0,
            ) if r.tiktok_blogger else None,
        )
        for r in rows
    ]
    return VideoSourceListResponse(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/download-all-zip")
async def download_all_zip_endpoint(
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    """Download all videos with a stored URL as a single zip file."""
    # Fetch all video sources (up to 1000)
    rows, _ = await list_video_sources(session, page=1, page_size=1000, owner_id=owner_id)
    videos = [(r, r.local_video_url or r.local_gcs_video_url) for r in rows]
    videos = [(r, url) for r, url in videos if url]

    async def generate_zip():
        # GCS URL 走 SDK（私有桶可用），其它走 httpx；统一通过 download_url_to_local 到临时文件
        import os as _os
        import tempfile
        from app.utils.gcs_download import download_url_to_local

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, mode="w", compression=zipfile.ZIP_STORED, allowZip64=True) as zf:
            for i, (v, url) in enumerate(videos, 1):
                tmp_fd, tmp_path = tempfile.mkstemp(suffix=".mp4")
                _os.close(tmp_fd)
                try:
                    await download_url_to_local(url, tmp_path, timeout=120.0)
                    safe_title = re.sub(r'[\\/*?:"<>|]', "_", v.video_title or v.blogger_name or "video")
                    filename = f"{i:03d}_{safe_title}.mp4"
                    with open(tmp_path, "rb") as f:
                        zf.writestr(filename, f.read())
                except Exception:
                    pass
                finally:
                    try:
                        _os.unlink(tmp_path)
                    except FileNotFoundError:
                        pass
        buf.seek(0)
        yield buf.read()

    return StreamingResponse(
        generate_zip(),
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="videos.zip"'},
    )


@router.get("/export-excel")
async def export_video_urls_excel(
    platform: str | None = Query(None),
    blogger_name: str | None = Query(None),
    tiktok_blogger_id: uuid.UUID | None = Query(None),
    tag_ids: str | None = Query(None),
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> Response:
    """导出当前过滤条件下视频的 TikTok 博主主页 URL、原始链接、CDN/GCS 存储链接为 Excel。"""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    parsed_tag_ids: list[uuid.UUID] = []
    if tag_ids:
        parsed_tag_ids = [uuid.UUID(part.strip()) for part in tag_ids.split(",") if part.strip()]

    rows, _ = await list_video_sources(
        session,
        page=1,
        page_size=10000,
        owner_id=owner_id,
        platform=platform,
        blogger_name=blogger_name,
        tiktok_blogger_id=tiktok_blogger_id,
        tag_ids=parsed_tag_ids,
    )

    # 导出前批量续签：拿到 Excel 的人短时间内（< 7 天）能直接点开播放
    from app.utils.gcs_signing import ensure_video_sources_signed_urls
    await ensure_video_sources_signed_urls(session, rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "视频链接导出"

    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(vertical="top", wrap_text=True)

    headers = ["TikTok 博主", "博主主页 URL", "视频原始链接", "local_video_url", "local_gcs_video_url"]
    ws.append(headers)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 60
    ws.row_dimensions[1].height = 22

    for i, r in enumerate(rows, start=2):
        blogger = r.tiktok_blogger
        ws.cell(row=i, column=1, value=(blogger.blogger_name if blogger else r.blogger_name) or "").alignment = center
        ws.cell(row=i, column=2, value=(blogger.blogger_url if blogger else "") or "").alignment = wrap
        ws.cell(row=i, column=3, value=r.source_url or "").alignment = wrap
        ws.cell(row=i, column=4, value=r.local_video_url or "").alignment = wrap
        ws.cell(row=i, column=5, value=r.local_gcs_video_url or "").alignment = wrap

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="video_urls.xlsx"'},
    )


@router.get("/{vs_id}", response_model=VideoSourceRead)
async def get_video_source_endpoint(
    vs_id: uuid.UUID,
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceRead:
    vs = await get_video_source_or_404(session, vs_id, owner_id)
    return await _vs_to_read(session, vs)


@router.get("/{vs_id}/stats-history", response_model=VideoSourceStatHistoryResponse)
async def get_stats_history_endpoint(
    vs_id: uuid.UUID,
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceStatHistoryResponse:
    """Get historical stats for a video source."""
    items = await get_stats_history(session, vs_id, owner_id)
    return VideoSourceStatHistoryResponse(items=items)  # type: ignore[arg-type]


@router.post("/{vs_id}/download", response_model=VideoSourceRead)
async def download_video_source_endpoint(
    vs_id: uuid.UUID,
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceRead:
    """Trigger background download via yt-dlp and upload to permanent storage."""
    vs = await trigger_download_and_upload(session, vs_id, owner_id)
    return await _vs_to_read(session, vs)


@router.patch("/{vs_id}/tags", response_model=VideoSourceRead)
async def replace_video_source_tags_endpoint(
    vs_id: uuid.UUID,
    payload: VideoSourceTagsUpdate,
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> VideoSourceRead:
    vs = await replace_video_source_tags(session, vs_id, payload.tag_ids, owner_id)
    return await _vs_to_read(session, vs)


@router.delete("/{vs_id}")
async def delete_video_source_endpoint(
    vs_id: uuid.UUID,
    owner_id: uuid.UUID | None = Depends(_get_owner_id),
    session: AsyncSession = Depends(get_db),
) -> Response:
    await delete_video_source(session, vs_id, owner_id)
    return Response(status_code=204)
